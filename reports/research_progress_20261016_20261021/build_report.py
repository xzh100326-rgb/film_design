"""Build a retrospective research record from frozen local experiment evidence.

Only writes inside this report directory. Never trains or changes experiments.
"""
from pathlib import Path
from datetime import datetime, timezone, timedelta
import hashlib
import json
import shutil

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

ROOT = Path(__file__).resolve().parent
WORK = Path(r'E:\lithography_machine_lens')
PROJECT = WORK / 'film_hppo'
ASSETS = ROOT / 'assets'
ASSETS.mkdir(parents=True, exist_ok=True)
OUTPUT = ROOT / '薄膜强化学习研究思考与实验记录_20261016至20261021.docx'

figures = {
    'duv': ('film_hppo/outputs/week_20261014_20261020/figures/duv_multiseed.png', 'fig01_duv_multiseed.png'),
    'masked': ('film_hppo_masked_diffusion/outputs/duv193_masked_hppo_seed1.png', 'fig02_masked_training.png'),
    'abs_search': ('film_hppo/outputs/week_20261014_20261020/figures/absorption_search_progress.png', 'fig03_absorption_search.png'),
    'abs_spectra': ('film_hppo/outputs/week_20261014_20261020/figures/spectra_comparison.png', 'fig04_duv_absorption_spectra.png'),
    'paper': ('film_hppo/outputs/references/oml_ppo_task2/table3_current_materials/comparison.png', 'fig05_paper42_recalculation.png'),
    'lamp_T': ('film_hppo/outputs/incandescent_improve/20261021_094617/spectra.png', 'fig06_lamp_T_comparison.png'),
    'hppo_training': ('film_hppo/outputs/hppo_spectral_search/20261021_145752/training_comparison.png', 'fig07_hppo_training.png'),
    'hppo_R': ('film_hppo/outputs/hppo_spectral_search/20261021_145752/reflection_comparison_dense.png', 'fig08_hppo_reflection.png'),
    'enhancement': ('film_hppo/outputs/hppo_spectral_search/20261021_145752/baseline_stage2/reflection.png', 'fig09_enhancement_reflection.png'),
}

designs = [
    ('D1', 'DUV 固定常数 v1 的最好设计', 'film_hppo/outputs/duv193__hppo__nk_v1__a00-45/20261018_160311/best_design.json',
     '193 nm；0–45°共10角度；2层；平均T=99.8246%；固定n且k=0。'),
    ('D2', 'DUV 固定常数 v2 的最好设计', 'film_hppo/outputs/duv193__hppo__nk_v2__a00-45/20261019_143507/best_design.json',
     '193 nm；0–45°共10角度；5层；平均T=99.8231%；不同于D1的光学常数。'),
    ('D3', '随机掩码填充试验的最好已遇设计', 'film_hppo_masked_diffusion/runs/duv193_masked_hppo_seed1/designs/best_design.json',
     '6层；文件记录平均T=99.2787%；沿用该分支原五材料和原配置，不能与D1/D2直接比较算法优劣。'),
    ('D4', '宽带吸收 HPPO 基线的当前物理条件重算', 'film_hppo/outputs/abs400-2000__hppo__baseline/20261020_004725/best_design_current_physics.json',
     '400–2000 nm；14层；平均A=93.1043%；用于和D5统一基底比较。'),
    ('D5', '宽带吸收联合搜索的最好设计', 'film_hppo/outputs/abs400-2000__hybrid__search/20261020_155601/best_design.json',
     '400–2000 nm；14层；总厚度600 nm；平均A=98.9035%；结构搜索与厚度优化结果。'),
    ('D6', '灯泡滤光片 T目标参数试验最好设计', 'film_hppo/outputs/incandescent_hppo_experiments/20261021_004517/best_training_design.json',
     '可见T=1；1层；score=0.62676255；可见T=77.98%，IR R=47.17%。'),
    ('D7', '灯泡滤光片 T目标联合搜索最好设计', 'film_hppo/outputs/incandescent_improve/20261021_094617/best_design.json',
     '可见T=1；16层；score=0.72059530；可见T=74.06%，IR R=72.47%。'),
    ('D8', '灯泡滤光片 R目标纯 HPPO 最好设计', 'film_hppo/outputs/hppo_spectral_search/20261021_145752/best_design.json',
     '可见R=0；80层；score=0.73244667；原网格MAE=0.26755333；总厚度24.9969 μm。'),
    ('D9', 'OML PPO 论文表3的42层参考结构', 'film_hppo/outputs/references/oml_ppo_task2/table3_current_materials/input_design.json',
     '论文结构；42层；总厚度8.54 μm；只作当前光学常数下的参考重算。'),
]

reports = [
    ('R1', 'DUV 与宽带吸收阶段汇总', 'film_hppo/outputs/week_20261014_20261020/README.md'),
    ('R2', '400 至 2000 nm 改进方法与限制', 'film_hppo/CHANGELOG.md'),
    ('R3', '灯泡 T目标12次 HPPO 参数试验', 'film_hppo/outputs/incandescent_hppo_experiments/20261021_004517/RESULTS.md'),
    ('R4', '灯泡 T目标联合搜索结果', 'film_hppo/outputs/incandescent_improve/20261021_094617/RESULTS.md'),
    ('R5', '灯泡 R目标纯 HPPO 完整结果', 'film_hppo/outputs/hppo_spectral_search/20261021_145752/RESULTS.md'),
    ('R6', '增强因子与光谱误差差异诊断', 'film_hppo/outputs/hppo_spectral_search/20261021_145752/ENHANCEMENT_GAP.md'),
    ('R7', '材料数据转换与来源', 'film_hppo/outputs/references/oml_ppo/README.md'),
    ('R8', '随机掩码填充分支说明', 'film_hppo_masked_diffusion/README.md'),
]

manifest = {'snapshot_at': datetime.now(timezone(timedelta(hours=8))).isoformat(timespec='seconds'),
            'scope': '2026-10-16 evening through 2026-10-21',
            'files': [], 'designs': []}
for key, (relative, filename) in figures.items():
    src = WORK / relative
    if not src.is_file():
        raise FileNotFoundError(src)
    shutil.copy2(src, ASSETS / filename)
    manifest['files'].append({'kind': 'figure', 'id': key, 'source': str(src),
                              'sha256': hashlib.sha256(src.read_bytes()).hexdigest()})
for ident, title, relative, note in designs:
    p = WORK / relative
    raw = p.read_bytes()
    data = json.loads(raw.decode('utf-8-sig'))
    manifest['designs'].append({'id': ident, 'title': title, 'path': str(p), 'note': note,
                               'sha256': hashlib.sha256(raw).hexdigest(),
                               'layer_count': data.get('layer_count'), 'score': data.get('score'),
                               'mean_transmission': data.get('mean_transmission'),
                               'mean_absorption': data.get('mean_absorption')})
for ident, title, relative in reports:
    p = WORK / relative
    manifest['files'].append({'kind': 'report', 'id': ident, 'source': str(p),
                              'sha256': hashlib.sha256(p.read_bytes()).hexdigest()})
(ROOT / 'evidence_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')

doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(1.9)
sec.bottom_margin = Cm(1.8)
sec.left_margin = Cm(2.0)
sec.right_margin = Cm(2.0)
sec.header_distance = Cm(.7)
sec.footer_distance = Cm(.75)

def font_style(style, size, bold=False):
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.insert(0, rfonts)
    for k, v in [('ascii', 'Calibri'), ('hAnsi', 'Calibri'), ('eastAsia', 'Microsoft YaHei')]:
        rfonts.set(qn('w:' + k), v)

for name, size, bold in [('Normal', 10.5, False), ('Title', 23, True), ('Subtitle', 11, False),
                          ('Heading 1', 15, True), ('Heading 2', 12, True), ('Caption', 9, False)]:
    font_style(doc.styles[name], size, bold)
doc.styles['Normal'].paragraph_format.line_spacing = 1.23
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(8)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(7)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
doc.styles['Caption'].paragraph_format.space_after = Pt(7)

def p(text='', bold=False, style=None, keep=False):
    para = doc.add_paragraph(style=style)
    para.add_run(text).bold = bold
    para.paragraph_format.keep_with_next = keep
    para.paragraph_format.widow_control = True
    return para

def h(text, level=1):
    return doc.add_heading(text, level=level)

def page():
    doc.add_page_break()

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    for cell, head, width in zip(t.rows[0].cells, headers, widths):
        cell.width = Cm(width); cell.text = head
    for row in rows:
        cells = t.add_row().cells
        for cell, val, width in zip(cells, row, widths):
            cell.width = Cm(width); cell.text = str(val)
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement('w:' + side); b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4'); b.set(qn('w:color'), 'D9D9D9'); borders.append(b)
    t._tbl.tblPr.append(borders)
    for ri, row in enumerate(t.rows):
        if ri == 0:
            repeat = OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(repeat)
        cant = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
        for ci, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side, value in [('top', '85'), ('bottom', '85'), ('left', '100'), ('right', '100')]:
                m = OxmlElement('w:' + side); m.set(qn('w:w'), value); m.set(qn('w:type'), 'dxa'); margins.append(m)
            tcpr.append(margins)
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'DFEAF3' if ri == 0 else ('F6F7F8' if ri % 2 == 0 else 'FFFFFF')); tcpr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.line_spacing = 1.14
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 or len(cell.text) > 14 else WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size = Pt(9)
                    run.bold = ri == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def link(label, path, size=9, show_path=False):
    path = Path(path)
    if not path.exists(): raise FileNotFoundError(path)
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.line_spacing = 1.08
    hyperlink = OxmlElement('w:hyperlink')
    rel = para.part.relate_to(path.as_uri(), RT.HYPERLINK, is_external=True)
    hyperlink.set(qn('r:id'), rel)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts'); fonts.set(qn('w:ascii'), 'Calibri'); fonts.set(qn('w:eastAsia'), 'Microsoft YaHei'); props.append(fonts)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size * 2))); props.append(sz)
    color = OxmlElement('w:color'); color.set(qn('w:val'), '245A81'); props.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
    run.append(props)
    text = OxmlElement('w:t'); text.text = str(path) if show_path else label; run.append(text)
    hyperlink.append(run); para._p.append(hyperlink)
    return para

figure_counter = 0
def figure(key, caption, max_height=8.4):
    global figure_counter
    figure_counter += 1
    original, filename = figures[key]
    path = ASSETS / filename
    with Image.open(path) as im:
        w, hi = im.size
    width = min(17., max_height * w / hi)
    para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(3)
    para.paragraph_format.keep_with_next = True
    shape = para.add_run().add_picture(str(path), width=Cm(width))
    docpr = shape._inline.docPr
    docpr.set('descr', caption)
    p(f'图{figure_counter}  {caption}  原图地址见附录B。', style='Caption')

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
for part in ['第 ', 'PAGE', ' 页']:
    if part == 'PAGE':
        fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)
    else:
        footer.add_run(part).font.size = Pt(9)
doc.core_properties.title = '薄膜强化学习研究思考与实验记录'
doc.core_properties.subject = '2026年10月16日晚至10月21日 组会问题的追问与实验验证'
doc.core_properties.author = '研究记录'

# Page 1: starting questions and the research logic, not an accomplishment list.
p('薄膜强化学习研究思考与实验记录', style='Title')
p('2026年10月16日晚至10月21日', style='Subtitle')
p('这份记录围绕上次组会留下的五个问题，回顾我如何提出猜想、开展实验，再根据结果修改判断。现阶段已经形成材料读取、多波段目标配置和实验复核能力；同时也确认，得到一个高分膜系、证明HPPO有效，以及满足实际光学应用，是三个需要分别验证的问题。')
h('1 组会留下的问题及最初排序')
table(['组会问题', '优先级', '当时希望先解决的事'], [
    ['建立 baseline', '最高', '判断现有HPPO究竟带来了多少收益'],
    ['材料标准化与扩展', '高', '保证数据可靠，再分阶段扩大材料空间'],
    ['掩码扩散与可回改设计', '中期重点', '验证允许修改已有膜层是否有效'],
    ['指数奖励', '小规模穿插', '增强接近最优区域时的奖励区分能力'],
    ['联合光学设计', '后续', '膜层优化可靠后再接入真实光学系统'],
], [4.2, 2.5, 10.3])
p('最初我倾向于通过扩大材料空间、改变生成方式和放大奖励差异来提高搜索能力。继续追问后发现，这些改动都需要一个共同前提：材料模型、物理边界、目标函数和评价口径必须清楚，否则分数上涨不能说明方法改进。')
h('思考推进的主线', 2)
p('材料名称是否够多 → 光学常数是否可靠 → 不同实验能否公平比较 → 搜索不到好结构是否源于HPPO → 外部优化提高了什么 → 训练奖励是否对应真实需求。')
p('以下正文保留每一步的动机、实际行动和新的疑问。D1至D9为已核对存在的设计文件，附录A提供完整地址；图表原始文件和主要实验报告见附录B。全部数值为仿真结果。')
p('时间锚点：10月17日完成掩码填充试验及诊断；18日至19日开展两套DUV光学常数的多seed实验；20日推进宽带吸收联合优化和论文数据接入；21日开展灯泡滤光片baseline与纯HPPO对照。')

page()
h('2 从扩大材料空间追问到材料数据可靠性')
h('2 1 材料空间首先是光学常数与加工条件的选择', 2)
p('我起初问的是，能否把所有适合薄膜设计的材料都放进动作空间，以及材料数据库中是否缺少LaF₃、AlF₃、GdF₃等DUV材料。随后问题转向：同一种材料为什么有多份折射率数据，为什么有些数据不覆盖193 nm，插值能否保证正确。')
p('由此形成的认识是，材料名称本身不能唯一决定n和k。薄膜与块体、晶向、沉积温度、工艺和样品厚度可能对应不同的数据版本。插值只能在已给数据范围内提供数值估计；它既不能证明实验数据正确，也不能把缺少的DUV数据变成已知数据。')
h('2 2 用两套固定光学常数检查 DUV 设计表现', 2)
p('为了继续做193 nm任务，我先用四种材料的固定光学常数构建可控实验，并比较两套MgF₂和LaF₃取值。两套数据都采用k=0的无吸收假设，分别开展10个seed实验。它们的最好角度平均透射率相近，但最优膜层数不同。')
table(['版本', 'MgF₂的n', 'LaF₃的n', '最好平均T', '对应层数'], [
    ['固定常数v1', '1.420', '1.696', '99.8246%', '2'],
    ['固定常数v2', '1.440', '1.640', '99.8231%', '5'],
], [3.6, 2.7, 2.7, 4.0, 3.0])
p('共同条件为193 nm、空气入射、CaF₂基底、0–45°共10个角度；AlF₃与GdF₃采用该阶段配置值。对应设计为D1、D2。不能把两个版本的差异直接归因于随机种子或算法。')
figure('duv', '两套固定光学常数的多seed结果。蓝线是最好已遇设计，橙线是最终评估均值，两者不能混用。', 7.0)

page()
h('2 3 把材料读取方式改成可追溯的数据选择', 2)
p('固定常数只适合193 nm单点，无法支持宽波段设计。因此我进一步提出，材料名称写进JSON后，让程序在refractiveindex.info数据库中查找材料；遇到多份YAML时由用户选择数据版本，同时保留逐波长n和k读取。')
p('材料接口逐步加入了数据版本选择、波段覆盖检查、来源与工艺说明、缺失k的显式处理以及文件哈希。没有AlF₃条目时保留材料需求并明确报错，而不是自动使用其他材料或猜测折射率。主项目随后移除了不再需要的lm读取路线。')
p('在查阅OML-PPO时，我又把作者仓库中的CSV和TXT转换为相同YAML格式，区分root、400–2000-acs、piano_black等来源版本。转换保留原始数据点，不把单位转换或插值当成新的实验测量。该项目仍采用线性插值，和原仓库的二次插值并非完全相同。')
h('这一阶段改变了我的判断', 2)
p('材料空间的扩展应当先保证每个条目的来源、工艺含义、波段覆盖及n和k完整性，再研究数量增加对搜索的影响。数据格式标准化已经推进；大材料空间下的系统稳定性对照仍未完成。DUV的高透射率也只是所选材料模型下的结果，不能直接解释为可加工器件已经达到同样性能。')
h('需要继续验证的问题', 2)
p('同一种算法在不同nk版本下是否仍稳定；扩材料后收益来自搜索能力还是引入了更合适的物理材料；图像提取和缺失吸收数据带来的不确定性有多大。这些问题要求后续实验固定数据版本并明确记录来源。')

h('3 从奖励区分能力追问到优化目标是否一致')
h('3 1 先区分终局奖励和每步奖励', 2)
p('面对episode很快收敛到一两步，我最初怀疑：前面加膜的奖励是0，是否会让模型只偏好EOS。进一步梳理GAE和价值网络后，我认识到，非终止步奖励为0并不等于这些动作得不到学习信号。γ=1且没有长度惩罚时，PPO的目标也不自动等价于最短路径。短膜局部最优、价值估计、探索和约束都可能影响终止倾向。')
p('因此我保留了不同奖励版本，用来区分“更快得到反馈”和“改变实际优化目标”这两件事。')
table(['奖励形式', '每步定义', '需要注意的目标差别'], [
    ['终局奖励', '完成膜系时r=F(D)，其他步为0', '直接对应最终膜系的评分'],
    ['每步绝对评分', 'rₜ=F(Dₜ)', '累计回报包含各中间膜系评分，可能偏好较长轨迹'],
    ['增量奖励', 'rₜ=g(Dₜ)−g(Dₜ₋₁)', 'γ=1且处理一致时可望远镜求和为最终与初始评分之差'],
], [3.3, 5.1, 8.6])

page()
h('3 2 组会中的指数奖励落地为对数混合评分', 2)
p('我希望放大0.998与0.999这类高性能结果的差别，实际实现采用归一化对数误差项，再和线性评分混合，而不是直接使用指数函数。')
p('MAE = mean(|S − S_target|)；score = 1 − MAE', bold=True)
p('log_score = −log10(max(MAE, ε)) / [−log10(ε)]', bold=True)
p('g = (1 − α) × score + α × log_score', bold=True)
p('α通过配置调节；增量版本返回相邻两步g的差。这样改善了高性能区域的数值区分，但没有消除长期信用分配和局部最优，也没有证明所有任务都因此受益。宽带吸收及后续灯泡主要实验又使用α=0的线性奖励，避免把奖励变换混入其他因素。')
h('3 3 我没有把膜层数增加当成奖励改进的充分证据', 2)
p('每步直接评分曾让模型探索更多层，但部分多层设计的最终光谱反而不如少层设计。由此我开始同时看最好设计、评估均值、完整光谱和实际物理目标，而不只看episode_steps或loss是否收敛。熵系数与多seed可以辅助探索，但目前没有证据表明一组参数能替代多seed验证。')

h('4 掩码填充试验推进了生成方式但还没有实现可回改')
p('基于组会建议，我提出预设20个膜层槽位，每步随机抽取若干空位，由HPPO一次输出这些位置的材料与厚度。TMM跳过MASK；出现EOS时，仅评价最左EOS之前的有效膜系。独立分支已实现并完成测试与训练。')
p('我继续检查了训练后为何更少选择EOS。128次诊断中，初始模型平均最终层数约5.80，训练后约18.41；平均最终评分由0.4927升至0.9297。这说明策略学会了不同的填充和终止行为，但不能据此证明它比顺序HPPO更好，因为材料与动作、奖励条件没有做严格匹配。')
p('更关键的是，当前版本只填尚未填的位置，没有重新掩码或修改已填好的层。因此组会最关心的“允许修改已有膜层是否有效”仍未被直接验证。当前成果应称为随机掩码填充HPPO试验，不能称为已经实现完整扩散模型。')

page()
figure('masked', '随机掩码填充版本的原始训练曲线。横轴为优化器更新次数；平均episode步数不等于最终物理膜层数。', 10.4)
p('该分支最好已遇设计D3的平均透射率约99.2787%，为6层。它使用该分支自己的五材料数据和约束；其数值与D1、D2不构成算法优劣对照。')
p('该分支也保存中间有效膜系，因此最好已遇设计与曲线中的评估终局均值不是同一指标。')
h('从这条分支得到的下一步问题', 2)
p('我需要检验的已从“能否换一种生成顺序”变成“同一材料数据、同一评价次数和同一约束下，能修改旧层的动作是否比只能加膜更有效”。这个问题适合后续单独设计对照；在当前主项目中，我继续优先检查数据与baseline。')

h('5 换到宽带吸收任务后暴露了搜索能力问题')
h('5 1 单点高透射并不能证明宽带任务同样容易', 2)
p('我把任务扩展到400–2000 nm高吸收，并问应该优化平均透射还是目标光谱。由此把环境评价扩展为R、T、A三种量，以及分波段目标和权重。单波段高吸收设置A=1；若只有某个波段有要求，其余波段不计分。')
p('纯HPPO的结果长期停在约92%至93%。我依次怀疑厚度限制、总厚度预算、基底、材料nk以及模型初始策略。连续厚度分支虽可输出很多小数位，但这只表示参数分辨率高，不能说明联合搜索已经找到了更好的厚度组合。')

page()
h('5 2 先核对物理条件再检查策略更新', 2)
p('这一阶段复查了材料来源和波长覆盖、论文近似nk、基底和背面边界，并比较了论文膜系与当前程序的TMM结果。基底后来按对应实验设置为n=1.45、500 nm有限玻璃、背面空气。比较设计时，需要用同一物理条件重新计算，不能沿用原文件中的旧评分。')
p('模型侧检查发现，材料选择可能过早集中，例如第一层选择TiO₂的概率显著升高。随后把连续PPO概率项改为只使用所选材料的厚度参数，并对EOS忽略厚度项。这修正了物理动作与损失的对应关系，但实验效果没有因此稳定改善。')
p('当我质疑“既然厚度已是连续的，为什么还要精调”时，问题实质变成：HPPO没有找到高性能结构，是搜索和训练困难，还是当前材料与约束根本不允许更好结果。')
h('5 3 用联合搜索检验当前空间是否存在更好解', 2)
p('为回答这个问题，引入improve：从历史HPPO候选出发，尝试材料插入、删除、替换；每次结构变化后，用TMM自动微分和SLSQP联合调整所有层厚度。厚度优化变量归一化，并处理单层边界与总厚度约束。')
figure('abs_search', '宽带吸收联合搜索的进展及最后14层厚度。提升来自结构邻域搜索和厚度优化，不是纯HPPO训练曲线。', 7.3)
p('在统一的当前基底下，原HPPO候选D4的平均吸收率为93.1043%；联合搜索D5达到98.9035%，14层，总厚度600 nm。它说明当前材料和约束下存在更好的解，但没有单独证明HPPO的初始材料顺序优于随机顺序。')

page()
figure('abs_spectra', '左图区分DUV的不同nk版本；右图在当前相同玻璃基底下比较HPPO与联合搜索吸收光谱。', 7.3)
h('5 4 结果提升之后我开始追问方法归因', 2)
p('我进一步问：随便给出一个材料顺序，是否也可能通过厚度优化达到相同结果；既然已经得到优良结构，为什么还需要训练VAC。由此区分了“找到一个好膜系”和“证明HPPO提供了有价值的搜索起点”。')
p('已有程序让VAC监督学习优良轨迹，并支持从该权重继续原生PPO训练。这是一条可用的工程路线，能够验证网络能否生成和保持高分设计，但不能把外部结构搜索、梯度优化与模仿预训练的收益全部归给纯HPPO。严格比较HPPO起点与随机起点，在同一后处理预算下谁更有效，仍需要专门实验。')
h('这一步形成的方法边界', 2)
p('吸收率98.9035%是全波段平均值，不代表每个波长都超过95%，也不代表加工后器件性能。材料中仍有从论文曲线近似提取的数据。后续报告应分别标记纯HPPO、联合优化、预训练VAC和更新后策略，保留各自的配置与结果来源。')
p('我由此明确了自己想要的通用性：每个新任务可以重新训练，但希望主体框架能够沿用；任务变化主要通过材料、边界和目标光谱配置表达，而不是每来一个任务就增加新的优化模块。')

page()
h('6 用灯泡滤光片建立更有辨别力的 baseline')
h('6 1 从单一平均指标转向完整目标光谱', 2)
p('为了检查框架在更复杂目标上的表现，我选择OML-PPO中的incandescent light bulb filter。它要求在可见波段低反射，同时在其余波段高反射。配置使用300–3000 nm、0–89°共90角度和s/p功率平均；10 nm波长采样是本项目的实现选择，不把它冒称为论文已公开的原始步长。')
p('材料采用作者仓库中对应版本数据，保留SiN与仓库Si₃N₄数据映射的说明。论文表3给出了42层、总厚度8.54 μm的结构，我将它作为固定参考，在当前光学常数和TMM下重新计算，而不是直接把论文图上的分数当作本程序的baseline。')
figure('paper', '论文42层参考结构的当前数据重算。包含R/T/A、角度依赖及诊断对照；改变k或偏振的曲线仅用于分析，不是最终配置。', 11.8)

page()
h('6 2 低反射与高透射的区别改变了实验问题', 2)
p('观察初始设计后，我发现红外反射较高时，可见光也可能被大量反射。进一步追问得到更重要的区别：若材料吸收不可忽略，可见R=0并不等于T=1。因此做了一轮以可见T=1为目标的独立诊断，和反射目标分开记录。')
table(['实验', '可见平均T', '红外平均R', '层数'], [
    ['纯HPPO的T目标最好设计 D6', '77.98%', '47.17%', '1'],
    ['提高红外训练权重的候选', '43.22%', '74.66%', '12'],
    ['T目标联合搜索设计 D7', '74.06%', '72.47%', '16'],
], [8.1, 3.1, 3.1, 2.7])
p('T目标诊断共12次训练，比较epoch=5或10、不同seed、熵系数、红外权重、最小层数和学习率。在统一1:4:3复核评分下，纯HPPO最好仍约0.62676，说明这些小规模参数调整没有解决任务。提高红外权重能改善IR R，却牺牲可见T；只看一个总分会掩盖这种取舍。')
figure('lamp_T', '可见T=1目标下的纯HPPO与16层联合搜索R/T/A对照。它与后续可见R=0的80层实验采用不同目标，分数不能直接排名。', 11.6)

page()
h('6 3 把增强因子加入评价并检查它的适用范围', 2)
p('我不再只问某个score是否“足够好”，而是要求用反射率曲线和论文中的增强因子χ共同评价。项目实现了固定输入功率下的反射率热辐射baseline，加入人眼响应、半球角度权重、波段外闭合假设和分辨率检查。χ是独立评价指标，没有自动成为训练奖励。')
table(['结构或候选', 'χ 当f为1', 'χ 当f为0.95', '方法'], [
    ['T目标的一层HPPO设计', '2.8418', '2.5044', '纯HPPO'],
    ['T目标16层设计', '5.6274', '4.4773', '联合优化'],
    ['R目标本次80层设计', '9.3264', '6.6558', '纯HPPO'],
    ['论文42层 当前nk重算', '14.3098', '9.2817', '固定参考结构'],
], [6.7, 3.0, 3.4, 3.9])
p('论文报告的16.60和10.67是文献值；当前重算没有严格复现它们。面积归一化、波段外光谱及原始实现细节存在限制，应把本程序的χ作为公开假设下的一致比较量。')
p('继续核对公式后，我发现当前baseline只读取反射率，f=1时有效逸出因子为1−R=T+A。它没有单独求解滤光片吸收、再辐射与实际透射光。因此低可见反射但高吸收的设计也可能得到较高χ，不能把χ直接解释成实际透射亮度倍率。')
h('从 baseline 得到的新问题', 2)
p('一个设计可以在均匀波长平均MAE上更好，却在人眼加权和黑体辐射加权的χ上更差。当前诊断发现，710–1500 nm的反射不足和可见敏感区域的反射偏高都影响χ。评价需要分清：目标光谱匹配、辐射回收模型及实际透射应用，它们的权重和假设不同。')
p('因此建立baseline已经从“找一条论文曲线”深入为“固定材料和物理条件，保留论文结构重算，说明积分权重与假设，并同时查看可见T和A”。')

page()
h('7 回到纯 HPPO 检验长序列奖励如何影响早期动作')
h('7 1 保持框架调整 GAE 和训练节奏', 2)
p('在经历联合优化后，我再次要求尽量保持HPPO框架，通过训练本身减小光谱误差。这里沿用可见R=0的目标和1:4:3波段权重，使用已有7材料、15–400 nm单层范围及80层上限；80层不是本轮新增限制。')
p('运行代码仅将写死的GAE λ=0.95开放为配置并增加合法性检查。先比较λ=0.95与1.0，各运行seed0和seed1。随后从较好的λ=1检查点继续训练，降低学习率和熵系数、增大每批采样量，并减少对同一批数据的重复训练。')
table(['阶段', '学习率', '熵系数', 'n_sample', 'batch', 'epoch', '环境步预算'], [
    ['第一阶段', '5×10⁻⁵', '0.001', '3000', '128', '5', '180000'],
    ['继续训练', '3×10⁻⁵', '0.0001', '6000', '256', '3', '240000'],
], [2.8, 2.5, 2.3, 2.4, 1.8, 1.7, 3.5])
p('λ=1的动机是增强终局奖励对早期动作的直接影响。它仍有方差代价，也没有添加层数奖励。训练用了18个分箱中心角度降低计算量，所有保存候选按原90角度复核，并另做2 nm波长网格检查。')
figure('hppo_training', '原生HPPO的评估均值与episode步数。继续训练横轴累计父阶段环境步；曲线均值和最好设计分数是不同统计量。', 10.8)

page()
h('7 2 光谱误差下降但应用问题仍未解决', 2)
table(['指标', '原纯HPPO最好', '本次纯HPPO最好'], [
    ['90角度 10 nm网格 加权MAE', '0.31365958', '0.26755333'],
    ['90角度 2 nm网格 加权MAE', '0.31293772', '0.26522247'],
    ['原网格score', '0.68634042', '0.73244667'],
    ['物理层数', '65', '80'],
    ['480–700 nm 平均R', '45.97%', '39.74%'],
    ['710–3000 nm 平均R', '86.56%', '91.54%'],
], [7.1, 4.95, 4.95])
figure('hppo_R', '纯HPPO最好候选与原设计及论文图线对比。2 nm网格；论文线来自图中提取，非同条件训练结果。', 7.2)
p('原网格MAE下降14.70%，2 nm网格下降15.25%。最好结构D8有80层、总厚度24.9969 μm，来自HPPO直接保存的候选，没有后续厚度优化。最好模型链路为180596步初始训练加245928步继续训练，其余预算用于对照；五阶段总计971512环境步。')
p('但可见T仅4.81%，A为55.45%。这个结果说明纯HPPO通过已有参数调整可以改善当前R目标，却仍没有满足高透射灯泡滤光片的真实需求。它也提示我：80层或更高score都不能直接作为成功判据。')
p('后续128次独立随机推理中，第一阶段λ=1 seed0的平均score约0.6841，继续训练最终检查点约0.7131；最好单设计0.7324不能代替策略均值。收益同时包含λ、训练节奏和训练预算变化，不能全部归因于单一参数。')
p('本次结构为80层、约25 μm，论文参考为42层、8.54 μm，计算与制造成本并不相同，不能据此宣称算法公平超越原文。')

page()
h('8 五个组会问题现在推进到了哪里')
table(['最初问题', '已经完成或形成的认识', '尚未完成的验证'], [
    ['建立baseline', '论文结构重算 光谱曲线 多seed 误差与χ的口径检查', '相同预算的随机或传统优化对照 实际透射与热模型评价'],
    ['材料标准化与扩展', 'YAML读取 版本与来源记录 多波长n和k 数据覆盖检查', '大材料空间的稳定性与工艺适配性系统测试'],
    ['掩码扩散与可回改', '20槽位随机填充HPPO分支及EOS行为分析', '重新修改已填膜层的动作 同条件消融比较'],
    ['指数奖励', '线性与归一化对数混合 增量奖励 独立分支', '同任务同预算下的受控对照 证明高性能区域实际收益'],
    ['联合光学设计', '意识到膜系评价需要联系实际光谱和入射条件', '与真实镜头光线角谱 加工容差和系统指标联合优化'],
], [3.6, 6.7, 6.7])
h('已经收敛下来的研究问题', 2)
p('我现在更关心的是：固定材料数据与约束后，纯HPPO在不同目标光谱上能否稳定找到高质量设计，以及主要瓶颈究竟来自搜索、长期信用分配，还是目标函数和应用需求不一致。每个任务重新训练是可以接受的，但希望同一个主体框架能够通过配置表达任务差异。')
h('下一步可以直接执行的验证', 2)
p('第一，明确灯泡任务最终要优化的是反射曲线复现，还是可见透射与红外回收，分别保留R目标和T目标，避免把两个分数混排。第二，在固定材料和物理条件下，比较纯HPPO、随机结构及传统优化，并统一TMM评价次数或计算预算。第三，对奖励和GAE采用多seed受控实验，分别记录最好结构、策略均值和达到阈值的成本。')
p('材料空间的继续扩展应随上述对照推进；可回改动作需要补齐真正修改旧层的功能。联合光学系统属于后续阶段，当前尚不能写成已完成工作。')
h('本次记录的结果边界', 2)
p('正文采用已经完成并复核的阶段结果，其中纯HPPO以10月21日16时08分的正式汇总为准。17时后另启动的R目标improve对照单独管理，其候选不纳入这里的纯HPPO成果；最新状态以原实验目录记录为准。本文不覆盖原始训练和配置文件。')

page()
h('附录A 最好设计与参考结构的完整地址')
p('以下路径均已逐项检查存在，并设置了本地文件链接；在本机支持本地链接的Word阅读器中可以打开。D6与D7采用可见T目标，D8采用可见R目标，不能直接按score跨目标排序。')
for ident, title, relative, note in designs:
    para = p(f'{ident}  {title}', bold=True, keep=True)
    p(note)
    link('', WORK / relative, size=9, show_path=True)
    if ident == 'D5':
        page()
        h('附录A 最好设计与参考结构的完整地址', 1)

page()
h('附录B 效果图与实验报告地址')
p('效果图已经嵌入正文；下列地址便于打开原分辨率图像。本文编号顺序与正文图序一致。原图副本及本次核对的文件哈希保存在报告目录assets和evidence_manifest.json。')
for i, (key, (relative, filename)) in enumerate(figures.items(), 1):
    if key == 'enhancement':
        # This optional source is indexed as a supplemental image, not a body figure.
        label = '补充图  增强因子基线的反射率比较'
    else:
        label = f'图{i}  原始效果图'
    p(label, bold=True, keep=True)
    link('', WORK / relative, size=8.5, show_path=True)
page()
h('附录B 主要实验报告与运行入口')
for ident, title, relative in reports:
    p(f'{ident}  {title}', bold=True, keep=True)
    link('', WORK / relative, size=8.5, show_path=True)
p('主项目路径', bold=True, keep=True)
link('', PROJECT, show_path=True)
p('R目标纯HPPO第一阶段复跑配置', bold=True, keep=True)
link('', PROJECT / 'configs/incandescent_hppo_gae1.json', show_path=True)
p('命令：python -m film_rl train --config configs/incandescent_hppo_gae1.json')
p('此命令会重新训练，不是只查看现有最好设计，也不自动执行第二阶段。已完成实验的全部参数、父检查点和预算请查看R5对应目录下的plan.json与EXPERIMENT_LOG.md。')

# Remove inherited title borders/theme colors, keep headings uniformly black.
for name in ['Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    ppr = style.element.find(qn('w:pPr'))
    if ppr is not None:
        for node in list(ppr):
            if node.tag == qn('w:pBdr'): ppr.remove(node)
    for color in style.element.findall('.//' + qn('w:color')):
        for attr in ['themeColor', 'themeTint', 'themeShade']:
            color.attrib.pop(qn('w:' + attr), None)
        color.set(qn('w:val'), '000000')

doc.save(OUTPUT)
print(str(OUTPUT))
print('Embedded figures:', figure_counter, 'Design links:', len(designs))
print('Sources verified:', len(manifest['files']) + len(manifest['designs']))
