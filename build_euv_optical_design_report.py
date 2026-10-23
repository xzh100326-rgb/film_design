from pathlib import Path
from datetime import date

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"E:\lithography_machine_lens")
OUT = ROOT / "EUV光刻投影光学系统设计流程与关键难点.docx"
ASSET_DIR = ROOT / "document_assets"
FLOW_IMG = ASSET_DIR / "euv_optical_design_flow.png"

FONT_CN = "Microsoft YaHei"
FONT_LATIN = "Arial"
NAVY = "17365D"
BLUE = "D9EAF7"
PALE_BLUE = "EEF5FB"
TEAL = "DDF1ED"
PURPLE = "EAE4F4"
ORANGE = "F7E7D2"
LIGHT_GRAY = "F2F2F2"
BORDER = "D9D9D9"
TEXT = "1F1F1F"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name=FONT_CN, size=10.5, bold=False, color=TEXT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), FONT_LATIN)
    rfonts.set(qn("w:hAnsi"), FONT_LATIN)
    rfonts.set(qn("w:eastAsia"), name)


def set_para_format(paragraph, before=0, after=5, line=1.32, first_line=True):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if first_line:
        fmt.first_line_indent = Cm(0.74)


def add_body_paragraph(doc, text, citations=None, bold_lead=None):
    p = doc.add_paragraph()
    set_para_format(p)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    if citations:
        r = p.add_run(" " + citations)
        set_run_font(r, size=9.5, color="3F5E7A")
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.25
    for r in p.runs:
        set_run_font(r)
    if not p.runs:
        r = p.add_run(text)
        set_run_font(r)
    else:
        p.runs[0].text = text
        set_run_font(p.runs[0])
    return p


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), FONT_LATIN)
    rfonts.set(qn("w:hAnsi"), FONT_LATIN)
    rfonts.set(qn("w:eastAsia"), FONT_CN)
    rpr.append(rfonts)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color="666666")
    begin_run = OxmlElement("w:r")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    begin_run.append(fld_begin)
    instr_run = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    instr_run.append(instr)
    sep_run = OxmlElement("w:r")
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    sep_run.append(fld_sep)
    text_run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    text_run.append(t)
    end_run = OxmlElement("w:r")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    end_run.append(fld_end)
    paragraph._p.append(begin_run)
    paragraph._p.append(instr_run)
    paragraph._p.append(sep_run)
    paragraph._p.append(text_run)
    paragraph._p.append(end_run)
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=9, color="666666")


def make_flow_diagram(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    w, h = 1800, 900
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    font_path = r"C:\Windows\Fonts\msyh.ttc"
    font = ImageFont.truetype(font_path, 30)
    small = ImageFont.truetype(font_path, 25)
    title_font = ImageFont.truetype(font_path, 38)
    draw.text((w // 2, 35), "EUV投影光学系统设计闭环", anchor="ma", font=title_font, fill="#17365D")

    steps = [
        ("01", "需求与性能指标", "CD 视场 通量 Overlay"),
        ("02", "NA 倍率与视场", "分辨率 景深 扫描场"),
        ("03", "系统架构选择", "反射式 环形场 遮拦"),
        ("04", "近轴共轭与光阑", "物像 瞳面 远心度"),
        ("05", "初始结构搜索", "镜组 遮拦 入射角"),
        ("06", "面形与像差优化", "非球面 自由曲面"),
        ("07", "照明掩模物镜协同", "Pupil Mask SMO"),
        ("08", "多层膜联合设计", "振幅 相位 偏振"),
        ("09", "公差与补偿器", "偏心 倾斜 间隔"),
        ("10", "热结构与杂散光", "稳定性 洁净度"),
        ("11", "制造计量与装调", "面形 粗糙度 波前"),
        ("12", "工作波长验收闭环", "成像 通量 迭代回灌"),
    ]
    colors = ["#D9EAF7"] * 3 + ["#DDF1ED"] * 3 + ["#EAE4F4"] * 3 + ["#F7E7D2"] * 3
    cols, rows = 3, 4
    x0, y0 = 90, 120
    bw, bh = 470, 135
    gx, gy = 95, 60
    coords = []
    for i in range(rows):
        row_indices = list(range(i * cols, (i + 1) * cols))
        if i % 2 == 1:
            row_indices = list(reversed(row_indices))
        for j, idx in enumerate(row_indices):
            x = x0 + j * (bw + gx)
            y = y0 + i * (bh + gy)
            coords.append((idx, x, y))
    coords.sort(key=lambda t: t[0])

    for idx, x, y in coords:
        draw.rounded_rectangle((x, y, x + bw, y + bh), radius=18, fill=colors[idx], outline="#9AA9B8", width=3)
        n, main, sub = steps[idx]
        draw.text((x + 28, y + 28), n, font=font, fill="#17365D")
        draw.text((x + 100, y + 25), main, font=font, fill="#1F1F1F")
        draw.text((x + 100, y + 79), sub, font=small, fill="#555555")

    def arrow(a, b):
        ax, ay = a
        bx, by = b
        draw.line((ax, ay, bx, by), fill="#566573", width=6)
        import math
        ang = math.atan2(by - ay, bx - ax)
        length = 22
        for delta in (2.55, -2.55):
            ex = bx + length * math.cos(ang + delta)
            ey = by + length * math.sin(ang + delta)
            draw.line((bx, by, ex, ey), fill="#566573", width=6)

    for idx in range(11):
        _, x1, y1 = coords[idx]
        _, x2, y2 = coords[idx + 1]
        c1 = (x1 + bw / 2, y1 + bh / 2)
        c2 = (x2 + bw / 2, y2 + bh / 2)
        if abs(y1 - y2) < 5:
            if x2 > x1:
                arrow((x1 + bw + 10, c1[1]), (x2 - 10, c2[1]))
            else:
                arrow((x1 - 10, c1[1]), (x2 + bw + 10, c2[1]))
        else:
            arrow((c1[0], y1 + bh + 8), (c2[0], y2 - 8))

    # feedback loop
    _, x12, y12 = coords[11]
    _, x1, y1 = coords[0]
    feedback_x = 35
    draw.line((x12 - 10, y12 + bh / 2, feedback_x, y12 + bh / 2), fill="#A35C15", width=5)
    draw.line((feedback_x, y12 + bh / 2, feedback_x, y1 + bh / 2), fill="#A35C15", width=5)
    arrow((feedback_x, y1 + bh / 2), (x1 - 10, y1 + bh / 2))
    draw.text((45, h - 30), "验收结果回灌需求、处方、膜系和公差模型", font=small, fill="#A35C15")
    img.save(path, dpi=(240, 240))


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_CN
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)

    title = styles["Title"]
    title.font.name = FONT_CN
    title.font.size = Pt(23)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    title.paragraph_format.space_after = Pt(9)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (("Heading 1", 16, 14, 7), ("Heading 2", 13, 10, 5), ("Heading 3", 11, 7, 3)):
        st = styles[name]
        st.font.name = FONT_CN
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        st._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        st._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    if "Figure Caption" not in [s.name for s in styles]:
        cap = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        cap = styles["Figure Caption"]
    cap.font.name = FONT_CN
    cap.font.size = Pt(9)
    cap.font.color.rgb = RGBColor.from_string("555555")
    cap._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(7)


def add_table(doc, headers, rows, widths, font_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.width = Cm(widths[i])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, top=120, bottom=120)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, size=font_size, bold=True, color="FFFFFF")
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = Cm(widths[i])
            set_cell_margins(cells[i], top=95, bottom=95)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2 == 1:
                set_cell_shading(cells[i], PALE_BLUE)
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in (0, len(row) - 1) and len(row) <= 5 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def build_document():
    make_flow_diagram(FLOW_IMG)
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.15)
    section.right_margin = Cm(2.15)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    props = doc.core_properties
    props.title = "EUV光刻投影光学系统设计流程与关键难点"
    props.subject = "13.5 nm EUV投影光学系统设计方法和文献综述"
    props.keywords = "EUV; 光刻; 投影物镜; 光学设计; 多层膜; High NA"
    props.author = ""

    # Footer
    fp = section.footer.paragraphs[0]
    add_page_number(fp)

    # Title block
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("EUV光刻投影光学系统设计流程与关键难点")
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(3)
    r = subtitle.add_run("面向13.5 nm投影物镜与High NA系统的工程化设计综述")
    set_run_font(r, size=12, color="555555")
    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_p.paragraph_format.space_after = Pt(14)
    r = date_p.add_run("2026年9月")
    set_run_font(r, size=10, color="777777")

    doc.add_heading("摘要", level=1)
    add_body_paragraph(doc, "EUV投影光学系统的设计目标，不是单独获得一个低像差镜头处方，而是在分辨率、视场、数值孔径、系统通量、波前误差、畸变、远心度、遮拦、多层膜角度响应、制造和装调稳定性之间建立可闭环验证的工程解。当前公开路线从0.33 NA、等倍率缩小的六镜系统，发展到0.55 NA、中央遮拦和变倍率半场成像的High NA系统。设计流程因此必须将光学处方、照明瞳面、反射式掩模、多层膜、热结构、计量和装调放在同一优化框架内。")
    add_body_paragraph(doc, "本文给出从需求分解到工作波长验收的完整流程，并归纳九类关键难点及其主流解决方法。结论基于公开期刊论文和SPIE会议论文；公开资料可以支持方法和工程权衡的判断，但不包含量产扫描机的完整镜面坐标、面形系数、装调灵敏度矩阵和供应链工艺窗口。")

    doc.add_heading("1 适用范围与系统边界", level=1)
    add_body_paragraph(doc, "本文主要讨论13.5 nm EUV光刻机中从反射式掩模到晶圆的投影光学系统，并覆盖与投影成像直接耦合的照明瞳面、多层反射膜、掩模入射角、晶圆侧NA、扫描倍率和工作波长计量。光源等离子体、集光器、真空系统和工件台控制只在其形成光学输入或误差预算时涉及。")
    add_body_paragraph(doc, "与DUV折射式镜头不同，EUV投影物镜是全反射、多镜、离轴、环形视场系统。每面镜既提供光焦度，也参与光路折叠和像差平衡；镜面上的多层膜则同时决定反射振幅和反射相位。因此，物镜处方和膜系不能完全分开设计。[1][10][11]")

    doc.add_heading("2 光学系统设计流程", level=1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(FLOW_IMG), width=Cm(16.6))
    doc_pr = run._element.xpath(".//wp:docPr")
    if doc_pr:
        doc_pr[0].set("descr", "EUV投影光学系统从需求定义到工作波长验收的十二步闭环流程图")
    cap = doc.add_paragraph("图1 EUV投影光学系统设计闭环", style="Figure Caption")

    process_rows = [
        ("1", "需求与性能指标", "波长、目标CD、视场、吞吐量、overlay、EPE、掩模尺寸、晶圆扫描条件", "可验证的系统级指标和误差预算"),
        ("2", "NA、倍率与视场", "分辨率、景深、掩模入射角、扫描场尺寸", "0.33或0.55 NA；4×或4×/8×变倍率方案"),
        ("3", "系统架构选择", "镜面数、环形场、中央遮拦、全场或半场", "反射式投影架构和机械包络"),
        ("4", "近轴共轭与光阑", "物像共轭、瞳共轭、主光线和远心度", "一级光学布局和光阑位置"),
        ("5", "初始结构搜索", "遮拦、镜片碰撞、入射角、工作距离", "可优化且可制造的球面或低阶非球面起点"),
        ("6", "面形与像差优化", "全视场波前、畸变、MTF、远心度、面形偏离量", "非球面或自由曲面处方"),
        ("7", "照明、掩模与物镜协同", "瞳面填充、照明形状、掩模3D效应、偏振", "可打印性和工艺窗口优化结果"),
        ("8", "多层膜联合设计", "各镜AOI分布、振幅、相位、偏振、峰值波长", "每面镜的横向渐变或非周期膜系"),
        ("9", "公差与补偿器", "偏心、倾斜、间隔、曲率、面形、膜厚误差", "装调公差、敏感度矩阵和补偿策略"),
        ("10", "热结构与杂散光", "热漂移、动态载荷、污染、非成像光路", "热机械稳定性和杂散光控制方案"),
        ("11", "制造、计量与装调", "面形、粗糙度、镀膜前后变形、镜面坐标", "可交付镜面、安装基准和as-built模型"),
        ("12", "工作波长验收闭环", "EUV波前、成像、通量、均匀性和扫描性能", "验收结果回灌处方、膜系、公差和校准模型"),
    ]
    add_table(doc, ["阶段", "设计活动", "主要输入", "主要输出"], process_rows, [1.0, 3.2, 6.0, 6.5], font_size=8.4)

    doc.add_heading("2.1 需求分解与一级参数", level=2)
    add_body_paragraph(doc, "首先把制程需求转换为光学指标，包括目标CD、数值孔径、景深、视场尺寸、缩小倍率、掩模主光线角、晶圆侧远心度、最大畸变、波前误差、通量和允许遮拦。NA提升可以提高分辨率，但会同时减小景深、扩大镜面口径和AOI跨度，因此NA不是独立输入，而是架构权衡的起点。[6][7]")

    doc.add_heading("2.2 架构和初始结构", level=2)
    add_body_paragraph(doc, "根据视场、倍率和机械包络选择镜面数、环形视场、中央遮拦以及全场或半场方案。随后建立物像和瞳面共轭，通过近轴搜索、镜组分解或正反向实光线追迹寻找不存在机械遮挡、AOI合理且具有像差平衡潜力的初始结构。[1]-[5]")

    doc.add_heading("2.3 高阶优化与成像协同", level=2)
    add_body_paragraph(doc, "初始结构被逐步转换为非球面或自由曲面，并在全视场内联合优化波前、畸变、远心度、最大非球面偏离量和镜面口径。对High NA系统，还要将遮拦瞳、照明形状、掩模三维效应和偏振纳入可打印性分析，而不能只依赖几何光线指标。[6]-[9]")

    doc.add_heading("2.4 膜系、公差与工程闭环", level=2)
    add_body_paragraph(doc, "获得几何处方后，需要把每面镜全口径、全视场AOI分布传递给膜系设计。膜层的反射振幅和相位被重新写入光学模型，再进行波前和瞳面通量优化。随后加入曲率、间隔、偏心、倾斜、面形、膜厚、热漂移和装调误差，选择可测量的补偿器，最终形成as-built模型并在工作波长下验收。[10]-[19]")

    doc.add_heading("3 关键设计难点与主流解决方法", level=1)
    diff_rows = [
        ("1", "高NA、视场与通量冲突", "NA提高会扩大镜面、AOI跨度并减小景深；完整大视场难以维持", "中央遮拦、4×/8×变倍率、半场扫描、强化焦控", "[6]-[9]"),
        ("2", "可行拓扑与初始结构稀缺", "全反射离轴系统易遮拦、碰撞或AOI超限；可行参数域很小", "环形场、近轴解域搜索、镜组分解、正反向光追", "[1]-[5]"),
        ("3", "高维非线性优化和局部极值", "曲率、间距、遮拦、像差和可制造性强耦合", "解析约束加全局搜索，再用局部算法精修；保留结构可行性约束", "[3]-[5]"),
        ("4", "全视场亚纳米波前与畸变", "中心视场低像差不代表扫描场内都满足波前、远心和畸变要求", "极端非球面、XY自由曲面、全视场逐级优化和面形偏离约束", "[4][7][8]"),
        ("5", "遮拦瞳的波前表达与成像", "中央遮拦改变瞳函数、衍射传递和常规Zernike表征", "优化遮拦尺寸；使用适用于遮拦瞳的基函数和RMS定义", "[6][9]"),
        ("6", "多层膜振幅相位与AOI耦合", "膜层会引入离焦、像散、倾斜和变倍率；高NA下角度和偏振效应增强", "膜系进入光学模型；横向渐变、深度渐变或非周期膜；优化瞳面均匀性", "[10]-[14]"),
        ("7", "多镜通量的指数损失", "单镜约65%至70%的反射率经多次反射后快速衰减", "控制镜面数和AOI；界面工程；按系统积分通量而非单镜峰值优化", "[15]-[17]"),
        ("8", "镜面制造与计量接近原子尺度", "极端非球面、粗糙度和中高频误差均影响波前和散射", "确定性修形、真空计量、EUV工作波长检测、镀膜前后面形闭环", "[7][15][18]"),
        ("9", "膜应力、稳定性与装调公差", "镀膜应力使基底变形；界面扩散、氧化和热漂移会改变长期性能", "缓冲层和退火补偿、B4C阻挡层、Ru封帽、蒙特卡洛公差和可调补偿器", "[5][16][17][19]"),
    ]
    add_table(doc, ["序号", "关键难点", "工程影响", "主流解决方法", "文献"], diff_rows, [0.8, 3.5, 4.5, 5.7, 1.8], font_size=8.1)

    doc.add_heading("3.1 高NA系统架构", level=2)
    add_body_paragraph(doc, "0.55 NA High NA系统的公开主线，是以中央遮拦降低投影物镜内部多层膜镜面的角度负担，并采用扫描方向8×、另一方向4×的变倍率设计，从全场转向半场。该方案换取更高分辨率和可接受的通量，但同时增加遮拦瞳成像、焦面控制和掩模侧角度管理的复杂度。[6][7][9]")

    doc.add_heading("3.2 初始结构与自动优化", level=2)
    add_body_paragraph(doc, "公开学术路线从反射式环形场系统的可行解域分类，发展到镜组分解、正反向实光线追迹以及模拟退火、粒子群与SQP组合的混合优化。共同原则是先满足遮拦、共轭、远心、AOI和工作距离，再释放高阶面形自由度；否则优化器容易获得数学上低像差、工程上不可实现的局部解。[1]-[5]")

    doc.add_heading("3.3 物镜与多层膜联合优化", level=2)
    add_body_paragraph(doc, "EUV多层膜不能近似为固定反射率边界。其振幅和相位随波长、入射角、偏振及镜面位置变化，反射相位会直接形成波前像差。现有主流方法是根据各镜AOI分布设计横向渐变膜，并在AOI跨度更大时引入深度渐变或纵向非周期结构；最终目标是全瞳、全视场的系统通量和波前，而不是法向入射下的单点峰值反射率。[10]-[14]")

    doc.add_heading("3.4 制造计量和装调", level=2)
    add_body_paragraph(doc, "光学设计必须在处方阶段限制最大非球面偏离量、局部斜率、镜面尺寸和公差敏感度。镜面加工、镀膜和装夹后都需要更新as-built模型；对关键波前误差，应通过真空或工作波长计量完成闭环。实验研究已经表明，曲面EUV镜的多层膜沉积可把附加面形误差控制在0.1 nm RMS以下，但这依赖沉积轮廓控制、基底质量和专用计量链。[7][15][18]")

    doc.add_heading("4 各阶段建议交付物", level=1)
    deliver_rows = [
        ("系统需求评审", "指标树、成像与通量预算、接口条件、验收方法", "每项指标有单位、视场条件和测量方式"),
        ("架构评审", "镜面数、倍率、环形场、遮拦、光阑和机械包络", "无不可消除遮拦；AOI和口径具备膜系及制造余量"),
        ("处方评审", "全视场波前、畸变、远心度、MTF、面形和光线数据库", "不能只提交中心视场或名义处方"),
        ("膜系评审", "各镜AOI地图、反射率相位矩阵、瞳面均匀性和膜厚容差", "膜系性能已回灌成像模型"),
        ("公差评审", "敏感度、蒙特卡洛结果、补偿器、测量和装调顺序", "补偿量可测、可调且不产生新的不可观测误差"),
        ("制造评审", "面形/粗糙度频带规范、镀膜应力预算、计量不确定度", "制造与计量能力覆盖设计允许区间"),
        ("验收评审", "EUV波前、通量、成像、扫描和长期稳定性结果", "实测数据进入as-built模型并关闭误差预算"),
    ]
    add_table(doc, ["评审节点", "主要交付物", "放行条件"], deliver_rows, [3.3, 8.2, 5.8], font_size=8.7)

    doc.add_heading("5 结论", level=1)
    add_body_paragraph(doc, "当前EUV投影光学系统的主流设计流程，是一条从制程需求出发、以可制造初始结构为基础、以全视场波前和系统通量为核心、以多层膜和热机械公差为耦合约束、最终通过工作波长计量闭环的系统工程流程。最关键的三项能力分别是：高NA架构权衡，亚纳米级光学制造与计量，以及物镜处方和多层膜的联合优化。任何只完成名义几何处方、却未纳入膜层相位、制造频带、公差补偿和as-built验证的设计，都还不能视为工程化投影物镜设计。")

    doc.add_page_break()
    doc.add_heading("参考文献", level=1)
    refs = [
        ("[1] M. Bal, F. Bociort, and J. Braat, Analysis, search, and classification for reflective ring-field projection systems, Applied Optics, 42, 2301-2311, 2003. ", "https://doi.org/10.1364/AO.42.002301"),
        ("[2] F. Liu and Y. Li, Grouping design of eight-mirror projection objective for high-numerical aperture EUV lithography, Applied Optics, 52, 7137-7144, 2013. ", "https://doi.org/10.1364/AO.52.007137"),
        ("[3] Y. Liu, Y. Li, and Z. Cao, Design of anamorphic magnification high-numerical aperture objective for EUV lithography by curvatures combination method, Applied Optics, 55, 4917-4923, 2016. ", "https://doi.org/10.1364/AO.55.004917"),
        ("[4] X. Yan, Y. Li, Y. Li, L. Liu, and K. Liu, Design of an EUV lithography projection objective with a grouping design method through forward and reverse real ray tracing, Applied Optics, 61, 7449-7454, 2022. ", "https://doi.org/10.1364/AO.469274"),
        ("[5] M. Ai, S. Xian, S. Li, Z. Huang, S. Hu, J. Liu, and J. Zhou, Aberration-balanced initial structure design of EUV lithography objective systems via hybrid optimization, Applied Optics, 65, 4802-4810, 2026. ", "https://doi.org/10.1364/AO.589602"),
        ("[6] E. van Setten et al., High NA EUV lithography: Next step in EUV imaging, Proceedings of SPIE 10957, 2019. ", "https://doi.org/10.1117/12.2514952"),
        ("[7] C. Zahlten, P. Graeupner, J. van Schoot, P. Kurz, J. Stoeldraijer, and W. Kaiser, High-NA EUV lithography: pushing the limits, Proceedings of SPIE 11177, 2019. ", "https://doi.org/10.1117/12.2536469"),
        ("[8] M. Liu and Y. Li, Anamorphic objective design for EUV lithography at the 5-1 nm technology node, Applied Optics, 60, 7254-7258, 2021. ", "https://doi.org/10.1364/AO.428136"),
        ("[9] L. de Winter et al., Extreme ultraviolet scanner with high numerical aperture: obscuration and wavefront description, Journal of Micro/Nanopatterning, Materials, and Metrology, 21, 023801, 2022. ", "https://doi.org/10.1117/1.JMM.21.2.023801"),
        ("[10] C. Liang, M. Descour, J. Sasian, and S. Lerner, Multilayer-coating-induced aberrations in EUV lithography optics, Applied Optics, 40, 129-135, 2001. ", "https://doi.org/10.1364/AO.40.000129"),
        ("[11] M. Bal, M. Singh, and J. Braat, Optimization of multilayer reflectors for EUV lithography, Journal of Microlithography, Microfabrication, and Microsystems, 3, 537-544, 2004. ", "https://doi.org/10.1117/1.1793171"),
        ("[12] A. Yakshin, I. Kozhevnikov, E. Zoethout, E. Louis, and F. Bijkerk, Properties of broadband depth-graded multilayer mirrors for EUV optical systems, Optics Express, 18, 6957-6971, 2010. ", "https://doi.org/10.1364/OE.18.006957"),
        ("[13] X. Zhong et al., High efficiency graded multilayer coating design using least-square fitting for NA0.55 EUV lithography objective, IWAPS, 2022. ", "https://doi.org/10.1109/IWAPS57146.2022.9972249"),
        ("[14] S. Haque, S. Uji, and M. Toyoda, Designing multilayer mirror structures to achieve high pupil transmittance uniformity in EUV lithography, Optical Engineering, 65, 045103, 2026. ", "https://doi.org/10.1117/1.OE.65.4.045103"),
        ("[15] C. Montcalm et al., Atomic-precision multilayer coating of the first set of optics for an EUV lithography prototype system, Applied Optics, 41, 3262-3269, 2002. ", "https://doi.org/10.1364/AO.41.003262"),
        ("[16] S. Bajt et al., Improved reflectance and stability of Mo-Si multilayers, Optical Engineering, 41, 1797-1804, 2002. ", "https://doi.org/10.1117/1.1489426"),
        ("[17] P. Mirkarimi, Stress, reflectance, and temporal stability of sputter-deposited Mo/Si and Mo/Be multilayer films for EUV lithography, Optical Engineering, 38, 1246-1259, 1999. ", "https://doi.org/10.1117/1.602170"),
        ("[18] R. Soufli et al., Sub-diffraction-limited multilayer coatings for the 0.3 numerical aperture micro-exposure tool for EUV lithography, Applied Optics, 46, 3736-3746, 2007. ", "https://doi.org/10.1364/AO.46.003736"),
        ("[19] E. Louis, A. Yakshin, T. Tsarfati, and F. Bijkerk, Nanometer interface and materials control for multilayer EUV-optical applications, Progress in Surface Science, 86, 255-294, 2011. ", "https://doi.org/10.1016/J.PROGSURF.2011.08.001"),
    ]
    for text, url in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        set_run_font(r, size=8.8)
        add_hyperlink(p, "DOI", url)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
